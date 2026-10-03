"""Injection-safe document summarizer - Streamlit UI (owner: Duclas - M5).

Run:  streamlit run app.py
Tabs: Summarize | Test my document | Side-by-side | Attack suite | Report | Prompts
"""
from __future__ import annotations

import io
import os

import pandas as pd
import streamlit as st

from src import eval as ev
from src import inject, pipeline, prompts, scanner, ui
from src.llm import DEFAULT_MODELS, KEY_ENV, LLMConfig

st.set_page_config(page_title="Injection-safe summarizer", page_icon=":material/shield:", layout="wide")

CASES = ev.load_cases()
CASE_BY_LABEL = {f"{c['id']} - {c['title']}": c for c in CASES}
STATUS_STYLE = {"ok": st.success, "rejected": st.warning, "refused": st.warning, "invalid_output": st.error, "error": st.error}


# ---------- sidebar: model settings ----------
def model_picker(prefix: str, default_provider: str | None = None) -> LLMConfig:
    providers = list(DEFAULT_MODELS)
    env_provider = (default_provider or os.getenv("LLM_PROVIDER") or "gemini").lower()
    provider = st.selectbox("Provider", providers, index=providers.index(env_provider) if env_provider in providers else 0,
                            key=f"{prefix}_provider")
    env_model = os.getenv("LLM_MODEL") if provider == env_provider else ""
    model = st.text_input("Model", value=env_model or DEFAULT_MODELS[provider], key=f"{prefix}_model_{provider}")
    has_env = bool(os.getenv(KEY_ENV[provider]))
    key = st.text_input("API key", type="password", key=f"{prefix}_key_{provider}",
                        placeholder="Loaded from .env" if has_env else f"Paste {KEY_ENV[provider]}")
    if not key and not has_env:
        st.caption(f":red[No key found for {provider}]")
    return LLMConfig(provider, model, key)


with st.sidebar:
    st.header("Model", divider="gray")
    CFG = model_picker("main")
    st.caption("Temperature is fixed at 0 for repeatable results.")
    st.header("Defence layers", divider="gray")
    st.markdown(
        "- **Gate**: rejects empty / non-document input\n"
        "- **L1 Scanner**: regex patterns, base64 decode, hidden HTML & zero-width removal\n"
        "- **L2 Spotlighting**: random boundary per request\n"
        "- **L3 Defence prompt**: rules, sandwich reminder, few-shot (v3+)\n"
        "- **L4 Output checks**: JSON schema + repair retry, refusal retry, leak & link redaction\n"
        "- **L5 Self-critique** (v3+): second prompt audits and rewrites"
    )


# ---------- helpers ----------
def read_upload(file) -> str:
    if file is None:
        return ""
    if file.name.lower().endswith(".pdf"):
        from pypdf import PdfReader
        return "\n".join(page.extract_text() or "" for page in PdfReader(io.BytesIO(file.getvalue())).pages)
    if file.name.lower().endswith(".docx"):
        import docx
        d = docx.Document(io.BytesIO(file.getvalue()))
        parts = [p.text for p in d.paragraphs if p.text.strip()]
        for table in d.tables:
            for row in table.rows:
                parts.append(" | ".join(c.text.strip() for c in row.cells if c.text.strip()))
        return "\n".join(parts)
    return file.getvalue().decode("utf-8", errors="replace")


def document_input(prefix: str) -> str:
    """Sample picker + upload + editable text area. Returns the document text."""
    def _load_sample():
        label = st.session_state[f"{prefix}_sample"]
        st.session_state[f"{prefix}_doc"] = CASE_BY_LABEL[label]["document"] if label in CASE_BY_LABEL else ""

    with st.container(horizontal=True):
        st.selectbox("Load a test case", ["(type your own)"] + list(CASE_BY_LABEL), key=f"{prefix}_sample",
                     on_change=_load_sample)
        up = st.file_uploader("Or upload a document", type=["txt", "md", "pdf", "docx"], key=f"{prefix}_upload")
    if up is not None and st.session_state.get(f"{prefix}_upname") != up.name:
        st.session_state[f"{prefix}_doc"] = read_upload(up)
        st.session_state[f"{prefix}_upname"] = up.name
    return st.text_area("Document", key=f"{prefix}_doc", height=220,
                        placeholder="Paste any document here - including one with hidden instructions.")


def show_result(res: pipeline.Result) -> None:
    STATUS_STYLE.get(res.status, st.info)(f"**{pipeline.STATUS_TEXT.get(res.status, res.status)}**"
                                          + (f" - {res.message}" if res.message and res.status != "ok" else ""))
    if res.summary:
        with st.container(border=True):
            st.markdown("**Summary**")
            st.write(res.summary)
    if res.version != "v1" and res.status == "ok":
        if res.injection_detected:
            st.error(f"Injection detected: {res.injection_notes or 'instructions in the document were ignored'}",
                     icon=":material/gpp_bad:")
        else:
            st.success("No injection detected", icon=":material/verified_user:")
    if res.scan is not None:
        st.caption(f"Scanner risk: **{res.scan.risk_level}** (score {res.scan.risk_score})  |  "
                   f"Model calls: {res.calls}  |  Latency: {res.latency_s}s  |  {res.model}")
    else:
        st.caption(f"Model calls: {res.calls}  |  Latency: {res.latency_s}s  |  {res.model}")
    if res.trace:
        ui.trace_chips(res.trace)
    with st.expander("Defence trace details", expanded=False):
        st.dataframe(pd.DataFrame(res.trace), hide_index=True)
        if res.scan and res.scan.findings:
            st.markdown("**Scanner findings**")
            st.dataframe(pd.DataFrame([vars(f) for f in res.scan.findings]), hide_index=True)
        if res.critic:
            st.markdown("**Self-critique verdict**")
            st.json(res.critic)
    with st.expander("Raw model output"):
        for i, r in enumerate(res.raw, 1):
            st.code(r or "(empty)", language="json" if r.strip().startswith("{") else None)


ui.hero(n_attacks=sum(c["kind"] == "attack" for c in CASES),
        n_types=len({c["type"] for c in CASES if c["kind"] == "attack"}), n_cases=len(CASES))

tab_sum, tab_test, tab_cmp, tab_suite, tab_rep, tab_prompts = st.tabs(
    [":material/summarize: Summarize", ":material/shield: Test my document", ":material/compare: Side-by-side", ":material/bug_report: Attack suite",
     ":material/assessment: Report", ":material/edit_note: Prompts"])

# ---------- Summarize ----------
with tab_sum:
    doc = document_input("sum")
    version = st.segmented_control("Prompt version", prompts.VERSIONS, default="v4", key="sum_version",
                                   format_func=lambda v: prompts.VERSION_LABELS[v])
    if st.button("Summarize", type="primary", icon=":material/play_arrow:", key="sum_run"):
        with st.spinner("Running defence pipeline..."):
            st.session_state.sum_res = pipeline.run(doc, version or "v4", CFG)
    if "sum_res" in st.session_state:
        show_result(st.session_state.sum_res)

# ---------- Test my document ----------
with tab_test:
    payloads = inject.load_payloads()
    st.caption(f"Give any clean document (e.g. your own documentation). The app plants each of {len(payloads)} attack "
               "types into it, one at a time, and checks whether each version still produces a clean summary. An attack "
               "counts as defended if none of its marker words reach the user.")
    doc_t = document_input("inj")
    with st.container(horizontal=True):
        t_versions = st.multiselect("Versions", prompts.VERSIONS, default=["v1", "v4"], key="inj_versions",
                                    format_func=prompts.VERSION_LABELS.get)
        where = st.radio("Where to plant the attack", inject.POSITIONS, horizontal=True, key="inj_where")
    t_types = st.multiselect("Attack types (empty = all)", [p["type"] for p in payloads], key="inj_types")
    st.caption(f"{len(t_types or payloads) * len(t_versions)} runs; v3 and v4 make 2 model calls per run.")
    if st.button("Run injection test", type="primary", icon=":material/play_arrow:", key="inj_run",
                 disabled=not t_versions):
        ok, reason = scanner.check_document(doc_t)
        if not ok:
            st.warning(reason)
        else:
            bar = st.progress(0.0, text="Starting...")
            st.session_state.inj_df = inject.run_injection_test(
                doc_t, t_versions, CFG, where, t_types or None, lambda d, t, s: bar.progress(d / t, text=f"{d}/{t}  {s}"))
            bar.empty()
    if "inj_df" in st.session_state:
        df = st.session_state.inj_df
        errors = df[df.status == "error"]
        if not errors.empty:
            st.error(f"{len(errors)} of {len(df)} runs got no answer from the model (API error), so they say nothing "
                     f"about the defence. Check the provider and API key in the sidebar.\n\n"
                     f"Error: {errors.iloc[0].message[:300]}", icon=":material/cloud_off:")
        summ = inject.summary_table(df)
        cols = st.columns(len(summ))
        for col, (_, row) in zip(cols, summ.iterrows()):
            with col:
                with st.container(border=True):
                    st.metric(prompts.VERSION_LABELS[row.version], f"{row.defence_rate}% defended")
                    st.caption(f"{int(row.defended)}/{int(row.attacks)} attacks defended  |  "
                               f"{int(row.flagged)} flagged as injection")
        st.dataframe(df[["version", "type", "defended", "status", "injection_detected", "markers_hit", "summary"]],
                     hide_index=True, column_config={"defended": st.column_config.CheckboxColumn("defended")})
        failed = df[~df.defended & (df.status != "error")]
        if failed.empty:
            st.success("Every planted attack was defended.", icon=":material/verified_user:")
        for _, r in failed.iterrows():
            with st.expander(f"Failed: {r.version} / {r.type} (leaked: {r.markers_hit or r.status})"):
                st.markdown("**Output shown to the user**")
                st.write(r.summary or "(empty)")
                st.markdown("**Document with the planted attack**")
                st.code(r.attacked_document, language=None, wrap_lines=True)
        st.download_button("Download CSV", df.drop(columns=["attacked_document"]).to_csv(index=False),
                           "injection_test.csv", "text/csv", icon=":material/download:", key="inj_dl")

# ---------- Side-by-side ----------
with tab_cmp:
    st.caption("Same input, two configurations. Compare prompt versions on one model, or two models on one version.")
    doc_c = document_input("cmp")
    mode = st.radio("Compare", ["Two prompt versions", "Two models"], horizontal=True, key="cmp_mode")
    left, right = st.columns(2)
    with left:
        with st.container(border=True):
            st.markdown("**Left**")
            lv = st.selectbox("Version", prompts.VERSIONS, index=0, key="cmp_lv", format_func=prompts.VERSION_LABELS.get)
            lcfg = CFG
            st.caption(f"Model: {CFG.provider}/{CFG.model}")
    with right:
        with st.container(border=True):
            st.markdown("**Right**")
            if mode == "Two prompt versions":
                rv = st.selectbox("Version", prompts.VERSIONS, index=3, key="cmp_rv", format_func=prompts.VERSION_LABELS.get)
                rcfg = CFG
                st.caption(f"Model: {CFG.provider}/{CFG.model}")
            else:
                rv = lv
                st.caption(f"Version: {prompts.VERSION_LABELS[lv]}")
                rcfg = model_picker("right")
    if st.button("Run both", type="primary", icon=":material/play_arrow:", key="cmp_run"):
        with st.spinner("Running both configurations..."):
            st.session_state.cmp_res = (pipeline.run(doc_c, lv, lcfg), pipeline.run(doc_c, rv, rcfg))
    if "cmp_res" in st.session_state:
        a, b = st.session_state.cmp_res
        c1, c2 = st.columns(2)
        with c1:
            st.subheader(f"{a.version} - {a.model}")
            show_result(a)
        with c2:
            st.subheader(f"{b.version} - {b.model}")
            show_result(b)

# ---------- Attack suite ----------
with tab_suite:
    st.caption(f"{len(CASES)} labelled cases: {sum(c['kind'] == 'attack' for c in CASES)} attacks across "
               f"{len({c['type'] for c in CASES if c['kind'] == 'attack'})} types, "
               f"{sum(c['kind'] == 'benign' for c in CASES)} benign, {sum(c['kind'] == 'offtopic' for c in CASES)} off-topic.")
    with st.expander("View test set"):
        st.dataframe(pd.DataFrame(CASES)[["id", "kind", "type", "title", "success_markers", "key_facts"]], hide_index=True)
    versions = st.multiselect("Versions", prompts.VERSIONS, default=["v1", "v4"], key="suite_versions")
    ids = st.multiselect("Cases (empty = all)", [c["id"] for c in CASES], key="suite_ids")
    delay = st.slider("Delay between calls (s) - raise this on free-tier rate limits", 0.0, 10.0, 0.0, 0.5)
    if st.button("Run attack suite", type="primary", icon=":material/play_arrow:", disabled=not versions):
        chosen = [c for c in CASES if not ids or c["id"] in ids]
        bar = st.progress(0.0, text="Starting...")
        df = ev.run_suite(versions, CFG, chosen, delay, lambda d, t, s: bar.progress(d / t, text=f"{d}/{t}  {s}"))
        bar.empty()
        st.session_state.suite_df = df
        st.session_state.suite_path = ev.write_report(df)
    if "suite_df" in st.session_state:
        df = st.session_state.suite_df
        st.success(f"Report saved to reports/{st.session_state.suite_path.name}")
        m = ev.metrics(df)
        cols = st.columns(len(m))
        for col, (_, row) in zip(cols, m.iterrows()):
            with col:
                with st.container(border=True):
                    st.metric(f"{row.version} overall pass rate", f"{row.overall_pass_rate}%")
                    st.caption(f"Attack defence {row.attack_defence_rate}%  |  Benign {row.benign_accuracy}%  |  "
                               f"Off-topic {row.offtopic_rejection}%")
        st.dataframe(m, hide_index=True)
        pt = ev.per_type(df)
        if not pt.empty:
            st.markdown("**Attack defence rate by attack type (%)**")
            st.bar_chart(pt, stack=False, horizontal=True)
        st.markdown("**All results**")
        st.dataframe(df, hide_index=True,
                     column_config={"passed": st.column_config.CheckboxColumn("passed")})
        st.download_button("Download CSV", df.to_csv(index=False), "results.csv", "text/csv", icon=":material/download:")

# ---------- Report ----------
with tab_rep:
    latest = ev.REPORTS / "latest_report.md"
    if latest.exists():
        st.markdown(latest.read_text(encoding="utf-8"))
    else:
        st.info("No report yet. Run the attack suite (or `python -m src.eval`) to generate one.")

# ---------- Prompts ----------
with tab_prompts:
    st.caption("Edit a prompt, save, then re-run it in the Summarize tab. Commit the change to keep prompt history.")
    files = sorted(p.name for p in prompts.PROMPT_DIR.iterdir() if p.suffix in (".md", ".json"))
    name = st.selectbox("Prompt file", files, key="prompt_file")
    text = st.text_area("Content", prompts.load(name), height=420, key=f"prompt_text_{name}")
    if st.button("Save prompt", icon=":material/save:"):
        prompts.save(name, text)
        st.toast(f"Saved {name}")
